"""Parse product-import files (.xlsx / .docx / .csv) into normalized rows.

A row is a dict: name, name_en, sku, barcode, category_name, unit,
price_cents, cost_cents, stock_qty, _line (1-based source row for messages).
Parsing never touches the database; validity that depends on live data
(duplicate barcodes etc.) is annotated by the caller.
"""
import csv
import io
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

MAX_ROWS = 1000

HEADER_ALIASES = {
    "name": {"名称", "商品名称", "品名", "商品", "产品名称", "货物名称",
             "name", "product", "product name", "item", "item name"},
    "name_en": {"英文名", "英文名称", "name_en", "english name"},
    "price": {"价格", "售价", "销售价", "单价", "零售价", "会员价",
              "price", "sale price", "selling price", "unit price"},
    "cost": {"成本", "进价", "成本价", "进货价", "批发价",
             "cost", "cost price", "purchase price"},
    "stock": {"库存", "数量", "存货", "库存数量", "现有库存",
              "stock", "qty", "quantity", "stock qty", "balance"},
    "barcode": {"条码", "条形码", "条码号", "国际条码",
                "barcode", "bar code", "ean", "ean-13", "ean13"},
    "sku": {"货号", "编号", "商品编号", "商品货号",
            "sku", "code", "item code", "product code"},
    "category": {"分类", "类别", "商品分类", "category", "group", "type"},
    "unit": {"单位", "unit", "uom"},
}
TOTAL_WORDS = {"合计", "总计", "总数", "total", "sum", "grand total"}


def _norm_header(cell):
    return re.sub(r"[\s_\-]+", " ", str(cell or "").strip().lower())


_ALIAS_LOOKUP = {}
for _field, _names in HEADER_ALIASES.items():
    for _n in _names:
        _ALIAS_LOOKUP[_norm_header(_n)] = _field


def _money_to_cents(value):
    """Parse a money cell to integer cents; None when unparseable/empty."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        d = Decimal(str(value))
    else:
        s = str(value).strip()
        if not s:
            return None
        s = s.replace(",", "").replace(" ", "")
        s = re.sub(r"^(rm|myr|¥|￥|\$)", "", s, flags=re.I)
        s = re.sub(r"(rm|myr)$", "", s, flags=re.I)
        if not s:
            return None
        try:
            d = Decimal(s)
        except InvalidOperation:
            return None
    return int((d * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _num(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_from_matrix(matrix):
    """matrix: list of rows, each a list of raw cells. Returns (rows, errors, warnings)."""
    rows, errors, warnings = [], [], []
    colmap = None
    start = 0
    if matrix:
        first = [_norm_header(c) for c in matrix[0]]
        found = {i: _ALIAS_LOOKUP[c] for i, c in enumerate(first) if c in _ALIAS_LOOKUP}
        fields = set(found.values())
        if "name" in fields and len(fields) >= 2:
            colmap = {f: i for i, f in found.items()}
            start = 1
    if colmap is None:
        colmap = {"name": 0, "price": 1, "cost": 2, "stock": 3, "barcode": 4,
                  "category": 5, "unit": 6, "sku": 7}
        warnings.append("未识别到表头行，已按默认列顺序解析：名称、价格、成本、库存、条码、分类、单位、货号")
    for idx in range(start, len(matrix)):
        raw = matrix[idx]
        if raw is None or all((c is None or str(c).strip() == "") for c in raw):
            continue
        line = idx + 1

        def cell(field):
            i = colmap.get(field)
            if i is None or i >= len(raw):
                return None
            return raw[i]

        name = _text(cell("name"))
        if not name:
            if any(_text(c) for c in raw):
                errors.append(f"第 {line} 行：缺少商品名称，已跳过")
            continue
        if name.strip().lower() in TOTAL_WORDS:
            continue
        price = _money_to_cents(cell("price"))
        if price is None:
            errors.append(f"第 {line} 行（{name}）：价格无法识别，已跳过")
            continue
        if price < 0:
            errors.append(f"第 {line} 行（{name}）：价格不能为负，已跳过")
            continue
        cost = _money_to_cents(cell("cost"))
        if cost is None:
            cost = 0
        if cost < 0:
            errors.append(f"第 {line} 行（{name}）：成本不能为负，已跳过")
            continue
        stock = _num(cell("stock"))
        if stock is None:
            stock = 0.0
        if stock < 0:
            errors.append(f"第 {line} 行（{name}）：库存不能为负，已跳过")
            continue
        rows.append({
            "name": name,
            "name_en": _text(cell("name_en")),
            "sku": _text(cell("sku")),
            "barcode": _text(cell("barcode")).replace(" ", ""),
            "category_name": _text(cell("category")),
            "unit": _text(cell("unit")) or "pcs",
            "price_cents": price,
            "cost_cents": cost,
            "stock_qty": stock,
            "_line": line,
        })
        if len(rows) >= MAX_ROWS:
            warnings.append(f"文件行数过多，只处理前 {MAX_ROWS} 条商品")
            break
    return rows, errors, warnings


def _parse_xlsx(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    matrix = [list(r) for r in ws.iter_rows(values_only=True)]
    return _rows_from_matrix(matrix)


def _parse_docx(data):
    import docx
    doc = docx.Document(io.BytesIO(data))
    rows, errors, warnings = [], [], []
    if doc.tables:
        for t in doc.tables:
            matrix = [[c.text for c in row.cells] for row in t.rows]
            r, e, w = _rows_from_matrix(matrix)
            rows.extend(r)
            errors.extend(e)
            warnings.extend(w for w in w if "默认列顺序" not in w or not rows)
    else:
        matrix = []
        for p in doc.paragraphs:
            line = (p.text or "").strip()
            if not line:
                continue
            parts = re.split(r"\t|,|，|;|；|\s{2,}", line)
            matrix.append([x.strip() for x in parts if x.strip() != ""])
        r, e, w = _rows_from_matrix(matrix)
        rows.extend(r)
        errors.extend(e)
        warnings.extend(w)
    if not doc.tables:
        warnings.append("Word 文档里没找到表格，已按文本行解析")
    if len(rows) > MAX_ROWS:
        rows = rows[:MAX_ROWS]
        warnings.append(f"文件行数过多，只处理前 {MAX_ROWS} 条商品")
    return rows, errors, warnings


def _parse_csv(data):
    text = None
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别（请另存为 UTF-8）")
    reader = csv.reader(io.StringIO(text))
    return _rows_from_matrix([list(r) for r in reader])


def parse_product_file(data, filename):
    """Returns {"rows": [...], "errors": [...], "warnings": [...], "format": str}."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        rows, errors, warnings = _parse_xlsx(data)
        fmt = "xlsx"
    elif name.endswith(".docx"):
        rows, errors, warnings = _parse_docx(data)
        fmt = "docx"
    elif name.endswith(".csv"):
        rows, errors, warnings = _parse_csv(data)
        fmt = "csv"
    else:
        raise ValueError("只支持 .xlsx / .docx / .csv 文件")
    return {"rows": rows, "errors": errors, "warnings": warnings, "format": fmt}
